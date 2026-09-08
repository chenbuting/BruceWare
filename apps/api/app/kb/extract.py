"""从资料里抽出可供检索的纯文本。抽不到就返回空。"""

from __future__ import annotations

import csv
import io
from pathlib import Path


_MAX_CHARS = 20000
_SHEET_MAX_CHARS = 40000
_SHEET_ROWS = 500
_SHEET_SHEETS = 8

_TEXT_SUFFIXES = {".txt", ".md", ".log", ".json", ".yaml", ".yml", ".ini"}
_SHEET_SUFFIXES = {".csv", ".xlsx", ".xls"}


def is_sheet_name(file_name: str) -> bool:
    """是不是表格文件。"""

    return Path(file_name or "").suffix.lower() in _SHEET_SUFFIXES


def extract_search_text(file_name: str, data: bytes) -> str:
    """按后缀抽正文，失败返回空字符串。"""

    suffix = Path(file_name or "").suffix.lower()
    text = ""
    try:
        if suffix in _TEXT_SUFFIXES:
            text = _decode_text(data)
        elif suffix == ".csv":
            text = _from_csv(data)
        elif suffix == ".xlsx":
            text = _from_xlsx(data)
        elif suffix == ".xls":
            text = _from_xls(data)
        elif suffix == ".docx":
            text = _from_docx(data)
        elif suffix == ".pdf":
            text = _from_pdf(data)
    except Exception:
        text = ""
    limit = _SHEET_MAX_CHARS if suffix in _SHEET_SUFFIXES else _MAX_CHARS
    return " ".join(text.split())[:limit]


def _decode_text(data: bytes) -> str:
    """先按 UTF-8，不行再按国标。避免 CSV 按错编码变成乱码。"""

    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gb18030", errors="replace")


def _join_cells(cells: list[str]) -> str:
    return " | ".join(item for item in cells if item)


def _from_csv(data: bytes) -> str:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text))
    lines: list[str] = []
    for index, row in enumerate(reader):
        if index >= _SHEET_ROWS:
            break
        line = _join_cells([str(cell).strip() for cell in row])
        if line:
            lines.append(line)
    return "\n".join(lines) or text


def _from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in book.worksheets[:_SHEET_SHEETS]:
            if sheet.title:
                lines.append(str(sheet.title).strip())
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= _SHEET_ROWS:
                    break
                line = _join_cells([str(cell).strip() for cell in row if cell is not None])
                if line:
                    lines.append(line)
    finally:
        book.close()
    return "\n".join(lines)


def _from_xls(data: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    lines: list[str] = []
    for sheet in book.sheets()[:_SHEET_SHEETS]:
        if sheet.name:
            lines.append(str(sheet.name).strip())
        for index in range(min(sheet.nrows, _SHEET_ROWS)):
            line = _join_cells([str(sheet.cell_value(index, col)).strip() for col in range(sheet.ncols)])
            if line:
                lines.append(line)
    return "\n".join(lines)


def _from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:40]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)
